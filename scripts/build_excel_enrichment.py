"""
Converts Manel's manual enrichment test (HUBSPOT_ENRICHMENT_REVIEW.xlsx,
sheet "Todas las Empresas") into a compact JSON keyed by hs_object_id, so the
review app can propose from it directly -- no name-matching needed, the
Excel already carries the real HubSpot company IDs.

Run once (re-run if Manel produces a newer version of the Excel):
    python3 scripts/build_excel_enrichment.py /path/to/HUBSPOT_ENRICHMENT_REVIEW.xlsx
"""
import json
import sys

import openpyxl

DEFAULT_SRC = "/Users/leandrobraier/Downloads/HUBSPOT_ENRICHMENT_REVIEW.xlsx"
OUT = "/Users/leandrobraier/Desktop/manelmasia/Projectes/nsign/Projectes/Neteja-CRM-HubSpot/hubspot-enrichment-review/data/excel_enrichment_by_id.json"

SHEET = "Todas las Empresas"


def main():
    src = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_SRC
    wb = openpyxl.load_workbook(src, data_only=True)
    ws = wb[SHEET]
    header = [c.value for c in ws[1]]
    idx = {name: i for i, name in enumerate(header)}

    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        obj_id = row[idx["hs_object_id"]]
        if not obj_id:
            continue
        out[str(int(obj_id))] = {
            "name": row[idx["name"]],
            "domain": row[idx["domain"]],
            "industry": row[idx["industry"]],
            "industry_sector_excel": row[idx["industry_sector"]],
            "company_size": row[idx["company_size"]],
            "website": row[idx["website"]],
            "research_confidence": row[idx["research_confidence"]],
            "research_source": row[idx["research_source"]],
            "research_method": row[idx["research_method"]],
            "notes": row[idx["notes"]],
        }

    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, indent=1, sort_keys=True)
    print(f"Wrote {len(out)} companies to {OUT}")


if __name__ == "__main__":
    main()
